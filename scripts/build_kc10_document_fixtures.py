#!/usr/bin/env python3
"""Build deterministic, public-safe PDF and DOCX ingestion fixtures for KC-10."""

from __future__ import annotations

import argparse
import datetime as dt
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen.canvas import Canvas


FIXED_DATE = dt.datetime(2026, 8, 14, 12, 0, 0, tzinfo=dt.timezone.utc)


def set_run_font(run, *, size: float, bold: bool = False, color: str = "000000") -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), color)


def ensure_child(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def set_width(parent, tag: str, width_dxa: int) -> None:
    width = ensure_child(parent, tag)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(width_dxa))


def apply_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != 9360:
        raise ValueError("Fixture table widths must sum to 9360 DXA")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_properties = table._tbl.tblPr
    set_width(table_properties, "w:tblW", 9360)
    table_indent = ensure_child(table_properties, "w:tblInd")
    table_indent.set(qn("w:type"), "dxa")
    table_indent.set(qn("w:w"), "120")
    ensure_child(table_properties, "w:tblLayout").set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_column = OxmlElement("w:gridCol")
        grid_column.set(qn("w:w"), str(width))
        grid.append(grid_column)

    for column_index, width in enumerate(widths):
        table.columns[column_index].width = Twips(width)
    for row in table.rows:
        row.height = None
        for column_index, cell in enumerate(row.cells):
            width = widths[column_index]
            cell.width = Twips(width)
            cell_properties = cell._tc.get_or_add_tcPr()
            set_width(cell_properties, "w:tcW", width)
            margins = ensure_child(cell_properties, "w:tcMar")
            for side, margin_width in (
                ("top", 80),
                ("bottom", 80),
                ("start", 120),
                ("end", 120),
            ):
                margin = ensure_child(margins, f"w:{side}")
                margin.set(qn("w:type"), "dxa")
                margin.set(qn("w:w"), str(margin_width))


def configure_docx_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ):
        style = document.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def build_docx(path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_docx_styles(document)

    document.core_properties.title = "Synthetic Asset Evidence Memo"
    document.core_properties.subject = "Public-safe KC-10 ingestion fixture"
    document.core_properties.author = "Knowledge Copilot Contributors"
    document.core_properties.created = FIXED_DATE
    document.core_properties.modified = FIXED_DATE
    document.core_properties.last_printed = FIXED_DATE

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    set_run_font(header.add_run("SYNTHETIC EVIDENCE MEMO"), size=9, bold=True, color="666666")

    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(16)
    title.paragraph_format.space_after = Pt(4)
    set_run_font(title.add_run("Synthetic Asset Evidence Memo"), size=23, bold=True)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    set_run_font(
        subtitle.add_run("Deterministic fixture for narrative and table extraction"),
        size=14,
        color="555555",
    )

    document.add_heading("Operations Summary", level=1)
    document.add_paragraph(
        "The synthetic asset completed a 24-room renovation in 2020. This document "
        "contains invented figures and no real company, property, or person data."
    )

    document.add_heading("Selected Metrics", level=2)
    document.add_paragraph(
        "Management used the following public-safe sample metrics during planning."
    )

    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Metric", "2019", "2020", "Unit"]
    for cell, value in zip(table.rows[0].cells, headers, strict=True):
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_fill(cell, "F2F4F7")
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=10, bold=True)
    for values in (
        ("Occupancy", "78.0", "72.5", "%"),
        ("ADR", "118.00", "123.45", "USD"),
        ("RevPAR", "92.04", "89.50", "USD"),
    ):
        cells = table.add_row().cells
        for cell, value in zip(cells, values, strict=True):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                set_run_font(run, size=10)
    apply_table_geometry(table, [2880, 1800, 1800, 2880])

    document.add_heading("Review Note", level=2)
    document.add_paragraph(
        "All values are synthetic and exist solely to verify stable source lineage."
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="kc10-docx-") as temp_dir:
        raw_path = Path(temp_dir) / "raw.docx"
        document.save(raw_path)
        normalize_docx_archive(raw_path, path)


def normalize_docx_archive(source: Path, destination: Path) -> None:
    with zipfile.ZipFile(source, "r") as archive:
        entries = [(name, archive.read(name)) for name in sorted(archive.namelist())]
    with zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for name, data in entries:
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            archive.writestr(info, data)


def build_pdf(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    canvas = Canvas(str(path), pagesize=letter, invariant=1)
    width, height = letter

    canvas.setTitle("Synthetic PDF Evidence")
    canvas.setAuthor("Knowledge Copilot Contributors")
    canvas.setFont("Helvetica-Bold", 20)
    canvas.drawString(72, height - 84, "Synthetic PDF Evidence")
    canvas.setFont("Helvetica", 11)
    canvas.drawString(72, height - 118, "Public-safe narrative fixture for deterministic ingestion.")
    canvas.setFont("Helvetica-Bold", 14)
    canvas.drawString(72, height - 166, "Operations Summary")
    canvas.setFont("Helvetica", 11)
    canvas.drawString(72, height - 192, "The synthetic asset completed a 24-room renovation in 2020.")
    canvas.drawString(72, height - 210, "All facts and figures in this file are invented for software testing.")
    canvas.setFont("Helvetica", 9)
    canvas.drawRightString(width - 72, 54, "Page 1")
    canvas.showPage()

    canvas.setFont("Helvetica-Bold", 20)
    canvas.drawString(72, height - 84, "OCR Review Sample")
    canvas.setFont("Helvetica", 11)
    canvas.drawString(72, height - 124, "The line below intentionally imitates a damaged OCR text layer.")
    canvas.setFont("Courier", 13)
    canvas.drawString(72, height - 166, "R^V|P@R {2O2O} = $8g.S0 || 0CR~REVIEW")
    canvas.setFont("Helvetica", 9)
    canvas.drawRightString(width - 72, 54, "Page 2")
    canvas.save()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    build_pdf(output_dir / "sample-evidence.pdf")
    build_docx(output_dir / "sample-evidence.docx")


if __name__ == "__main__":
    main()
